from pathlib import Path
from datetime import date, datetime

import pandas as pd
from docx import Document
from pptx import Presentation
from openpyxl import load_workbook

OUTPUT_DIR = Path(__file__).resolve().parent / "music"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
GENERATED_DATE = date.today()
DISPLAY_DATE = GENERATED_DATE.strftime("%B %d, %Y").replace(" 0", " ")
GENERATED_AT = datetime.now()
DOCUMENT_TITLE = f"Music Session Transcript Insights — {DISPLAY_DATE}"
EXCEL_TITLE = f"Session Action Matrix — {DISPLAY_DATE}"
PRESENTATION_TITLE = f"Studio Creative Review — {DISPLAY_DATE}"

# Simulated text transcript data pulled from a session log
transcript_notes = [
    {"Timestamp": "00:15:20", "Category": "Lyrics", "Feedback": "Decided to change the second verse to focus on family conflict themes."},
    {"Timestamp": "00:28:45", "Category": "Vocal Delivery", "Feedback": "Tone needs to feel more vulnerable and conversational on the bridge."},
    {"Timestamp": "00:42:10", "Category": "Arrangement", "Feedback": "Agreed to strip back the acoustic guitar layer to let the vocals breathe."}
]

# 1. Word Document: Comprehensive Session Log
doc = Document()
doc.core_properties.title = DOCUMENT_TITLE
doc.core_properties.created = GENERATED_AT
doc.core_properties.modified = GENERATED_AT
doc.add_heading(DOCUMENT_TITLE, 0)
doc.add_paragraph(f"Generated: {DISPLAY_DATE}")
doc.add_paragraph("Extracted notes and decisions from spoken studio dialogue.")

for note in transcript_notes:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(f"[{note['Timestamp']}] {note['Category']}: ").bold = True
    p.add_run(note['Feedback'])

doc.save(OUTPUT_DIR / "Session_Transcript_Summary.docx")

# 2. Excel Spreadsheet: Categorized Action Tracking
df = pd.DataFrame(transcript_notes)
excel_path = OUTPUT_DIR / "Session_Action_Matrix.xlsx"
df.to_excel(excel_path, index=False)
workbook = load_workbook(excel_path)
workbook.properties.title = EXCEL_TITLE
workbook.properties.created = GENERATED_AT
workbook.properties.modified = GENERATED_AT
workbook.save(excel_path)

# 3. PowerPoint: Creative Direction Deck
prs = Presentation()
prs.core_properties.title = PRESENTATION_TITLE
prs.core_properties.created = GENERATED_AT
prs.core_properties.modified = GENERATED_AT
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = PRESENTATION_TITLE
slide.placeholders[1].text = f"Key Decisions from Spoken Transcripts — {DISPLAY_DATE}"

bullet_slide = prs.slides.add_slide(prs.slide_layouts[1])
bullet_slide.shapes.title.text = "Action Items & Flow Adjustments"
tf = bullet_slide.placeholders[1].text_frame
tf.text = "Highlights from Session Dialogue:"

for note in transcript_notes:
    p = tf.add_paragraph()
    p.text = f"{note['Category']}: {note['Feedback']}"

prs.save(OUTPUT_DIR / "Session_Creative_Direction.pptx")

print("Generated Office suite documents successfully using text transcripts!")
