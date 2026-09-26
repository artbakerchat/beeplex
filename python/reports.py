"""Office exports, loaded only when requested. Importing has no side effects."""

from datetime import date

from .config import DATA_DIR, DEMO


def generate(limit=10) -> dict:
    try:
        from docx import Document
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from pptx import Presentation
    except ImportError:
        from .client import BeeError

        raise BeeError(
            'Report dependencies are missing. Install with: python -m pip install "beeplex[reports]" (or pip install -e ".[reports]" from the repository).'
        ) from None

    from .bee_fetcher import fetch_report_data

    rows, info = fetch_report_data(limit=limit, persist=True)
    mode = "demo" if DEMO else "live"
    if not rows:
        return {
            "mode": mode,
            "count": 0,
            "files": [],
            "message": "No conversations found; no Office reports generated.",
        }
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    dates = [row["Recording_Date"] for row in rows if row.get("Recording_Date")]
    day = max(dates) if dates else date.today()
    label = f"Bee conversations — {day}" + (" [DEMO]" if DEMO else "")

    document = Document()
    document.add_heading(label, 0)
    document.add_paragraph(f"Source: {info['detail']}")
    for row in rows:
        document.add_heading(row["Session_Title"], 1)
        for key, value in row.items():
            if key != "Session_Title":
                document.add_paragraph(f"{key.replace('_', ' ')}: {value or '—'}")
    doc_path = DATA_DIR / f"Bee_Transcript_Action_Log_{day}.docx"
    document.save(doc_path)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Conversations"
    columns = list(rows[0])
    sheet.append([column.replace("_", " ") for column in columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append([str(row.get(key) or "") for key in columns])
        # Transcripts and titles are text, including strings starting with '='.
        for cell in sheet[sheet.max_row]:
            cell.data_type = "s"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 30
    xlsx_path = DATA_DIR / f"Bee_Transcript_Metrics_{day}.xlsx"
    workbook.save(xlsx_path)

    slides = Presentation()
    title = slides.slides.add_slide(slides.slide_layouts[0])
    title.shapes.title.text = label
    title.placeholders[1].text = f"{len(rows)} conversations · {info['detail']}"
    for row in rows:
        slide = slides.slides.add_slide(slides.slide_layouts[1])
        slide.shapes.title.text = row["Session_Title"]
        slide.placeholders[1].text = "\n".join(
            [
                f"Engagement: {row['Engagement_Level']}",
                f"Forward motion: {row['Forward_Motion']}",
                f"Tone: {row['Tone_Rating']}",
                f"Summary: {row['Summary_Notes']}",
                f"Actions: {row['Action_Items'] or 'None recorded'}",
            ]
        )
    pptx_path = DATA_DIR / f"Bee_Insights_Presentation_{day}.pptx"
    slides.save(pptx_path)
    return {
        "mode": mode,
        "count": len(rows),
        "detail": info["detail"],
        "files": [
            str(path)
            for path in (doc_path, xlsx_path, pptx_path, DATA_DIR / "dashboard.html")
        ],
    }


if __name__ == "__main__":
    import argparse
    import json

    parser = argparse.ArgumentParser(description="Generate Bee conversation reports.")
    parser.add_argument(
        "--limit", type=int, choices=range(1, 51), default=10, metavar="1..50"
    )
    args = parser.parse_args()
    print(json.dumps(generate(args.limit), indent=2))
