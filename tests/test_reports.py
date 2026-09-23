from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from beeplex import bee_fetcher, llm_scoring, reports
from beeplex.demo_cli import _conversations


def fixture_rows():
    """Deterministic rows scored from the bundled demo conversations."""
    return [
        bee_fetcher._row_from_source(
            conv,
            bee_fetcher._utterance_parts(conv),
            bee_fetcher._utterance_events(conv),
            llm=None,
            conv_id=conv["id"],
        )
        for conv in _conversations()[:2]
    ]


def test_exports_contain_requested_rows_and_preserve_formula_like_text(
    tmp_path, monkeypatch
):
    rows = fixture_rows()
    rows[0]["Session_Title"] = '=HYPERLINK("https://example.invalid", "untrusted")'
    requested = []

    def fetch(limit, *, persist):
        requested.append((limit, persist))
        return rows[:limit], {"mode": "demo", "detail": "test fixture"}

    monkeypatch.setattr(bee_fetcher, "fetch_report_data", fetch)
    monkeypatch.setattr(reports, "DATA_DIR", tmp_path)
    monkeypatch.setattr(reports, "DEMO", True)
    result = reports.generate(limit=2)
    assert requested == [(2, True)]
    workbook = load_workbook(result["files"][1])
    assert workbook.active.max_row == 3
    title_cell = workbook.active["B2"]
    assert title_cell.value == rows[0]["Session_Title"]
    assert title_cell.data_type == "s"
    workbook.close()
    assert len(Presentation(result["files"][2]).slides) == 3
    paragraphs = [p.text for p in Document(result["files"][0]).paragraphs]
    assert rows[0]["Session_Title"] in paragraphs
    assert rows[1]["Session_Title"] in paragraphs
    assert all(Path(path).is_file() for path in result["files"][:3])


def test_llm_requires_explicit_opt_in_even_with_credentials(monkeypatch):
    monkeypatch.setattr(llm_scoring, "LLM_ENABLED", False)
    monkeypatch.setattr(llm_scoring, "API_KEY", "configured-but-not-authorized")

    def fail_network(*args, **kwargs):
        raise AssertionError("Provider must not be contacted without opt-in")

    monkeypatch.setattr(llm_scoring.urllib.request, "urlopen", fail_network)
    assert not llm_scoring.available()
    assert (
        llm_scoring.llm_engagement_batch([("1", [("You", "Some private words")])]) == {}
    )
    assert llm_scoring.llm_text("private diary prompt") is None
