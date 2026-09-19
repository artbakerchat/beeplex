from pathlib import Path
from datetime import date, datetime

import pandas as pd
from docx import Document
from pptx import Presentation
from openpyxl import load_workbook

OUTPUT_DIR = Path(__file__).resolve().parent / "family"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
GENERATED_DATE = date.today()
DISPLAY_DATE = GENERATED_DATE.strftime("%B %d, %Y").replace(" 0", " ")
GENERATED_AT = datetime.now()
DOCUMENT_TITLE = f"Bee Wearable - Conversation & Transcript Log — {DISPLAY_DATE}"
EXCEL_TITLE = f"Bee Transcript Metrics — {DISPLAY_DATE}"
PRESENTATION_TITLE = f"Bee Wearable AI Insights — {DISPLAY_DATE}"

# 1. Enhanced dataset using clear, descriptive categories for Tone and Engagement
bee_transcript_data = [
    {
        "Session_Title": "Product Strategy Sync",
        "Source_Transcript_Snippet": "...we need to finalize the Q3 roadmap by Friday and assign module owners...",
        "Key_Topic": "Roadmap Deadlines",
        "Tone_Rating": "7/10 (Moderate Positive)",
        "Engagement_Level": "High",
        "Action_Items": "Finalize Q3 roadmap by Friday; assign module owners.",
        "Summary_Notes": "Team aligned on core priorities and established strict delivery milestones."
    },
    {
        "Session_Title": "Client Feedback Review",
        "Source_Transcript_Snippet": "...the client mentioned latency issues during peak hours, need an urgent patch...",
        "Key_Topic": "Performance Bug",
        "Tone_Rating": "4/10 (Low / Tense)",
        "Engagement_Level": "High",  # Notice High Engagement despite Low Tone (Urgent crisis)
        "Action_Items": "Deploy latency patch before Monday peak hours.",
        "Summary_Notes": "Addressed urgent customer friction point; engineering team to investigate."
    },
    {
        "Session_Title": "Weekly Team Catch-up",
        "Source_Transcript_Snippet": "...everyone's workload looks balanced, let's keep the current sprint velocity...",
        "Key_Topic": "Workload & Velocity",
        "Tone_Rating": "9/10 (High Positive)",
        "Engagement_Level": "Moderate",
        "Action_Items": "Maintain current sprint tasks; schedule next retro.",
        "Summary_Notes": "Positive alignment; team morale is high and pacing is sustainable."
    }
]

# 2. Generate Word Document
doc = Document()
doc.core_properties.title = DOCUMENT_TITLE
doc.core_properties.created = GENERATED_AT
doc.core_properties.modified = GENERATED_AT
doc.add_heading(DOCUMENT_TITLE, 0)
doc.add_paragraph(f"Generated: {DISPLAY_DATE}")
doc.add_paragraph("Processed from text transcripts (Raw audio is automatically deleted post-transcription by Bee).")

for item in bee_transcript_data:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(f"Session: {item['Session_Title']} — Topic: {item['Key_Topic']} ").bold = True
    p.add_run(f"(Tone: {item['Tone_Rating']} | Engagement: {item['Engagement_Level']})\n")
    p.add_run(f"Transcript Excerpt: \"{item['Source_Transcript_Snippet']}\"\n").italic = True
    p.add_run(f"Action Items: {item['Action_Items']}\n")
    p.add_run(f"Summary: {item['Summary_Notes']}")

doc.save(OUTPUT_DIR / "Bee_Transcript_Action_Log.docx")

# 3. Generate Excel Spreadsheet
df = pd.DataFrame(bee_transcript_data)
excel_path = OUTPUT_DIR / "Bee_Transcript_Metrics.xlsx"
df.to_excel(excel_path, index=False)
workbook = load_workbook(excel_path)
workbook.properties.title = EXCEL_TITLE
workbook.properties.created = GENERATED_AT
workbook.properties.modified = GENERATED_AT
workbook.save(excel_path)

# 4. Generate PowerPoint Presentation
prs = Presentation()
prs.core_properties.title = PRESENTATION_TITLE
prs.core_properties.created = GENERATED_AT
prs.core_properties.modified = GENERATED_AT
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = PRESENTATION_TITLE
slide.placeholders[1].text = f"Transcript Summaries & Action Items Review — {DISPLAY_DATE}"

bullet_slide = prs.slides.add_slide(prs.slide_layouts[1])
bullet_slide.shapes.title.text = "Key Takeaways from Transcripts"
tf = bullet_slide.placeholders[1].text_frame
tf.text = "Summary of Bee Device Transcripts:"

for item in bee_transcript_data:
    p = tf.add_paragraph()
    p.text = f"{item['Session_Title']} [{item['Tone_Rating']}] -> Action: {item['Action_Items']}"

prs.save(OUTPUT_DIR / "Bee_Insights_Presentation.pptx")

print("Successfully generated updated Bee Wearable Word, Excel, and PowerPoint files!")
